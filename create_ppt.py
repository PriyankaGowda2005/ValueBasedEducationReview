"""
Script to create PowerPoint presentation for Value-Based Education Literature Review
Requires: python-pptx library (install with: pip install python-pptx)
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_presentation():
    # Create presentation object
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    
    # Define slide layouts
    title_slide_layout = prs.slide_layouts[0]  # Title slide
    content_slide_layout = prs.slide_layouts[1]  # Title and content
    blank_layout = prs.slide_layouts[6]  # Blank layout
    
    # SLIDE 1: TITLE SLIDE
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = "Value-Based Education and Mental Health in School Students"
    subtitle.text = "A Narrative Literature Review with Scenario-Based Analytical Synthesis\n\nTraditional Indian Texts (Basavanna's Vachanas, Hitopadesha, Panchatantra)\nas Conceptual Frameworks\n\n[This is a narrative literature review - not an empirical study]"
    
    # SLIDE 2: DISCLOSURE STATEMENT
    slide = prs.slides.add_slide(content_slide_layout)
    title = slide.shapes.title
    content = slide.placeholders[1]
    
    title.text = "CRITICAL POSITIONING"
    tf = content.text_frame
    tf.text = "• This presentation presents a NARRATIVE LITERATURE REVIEW"
    p = tf.add_paragraph()
    p.text = "• NOT an empirical intervention study"
    p = tf.add_paragraph()
    p.text = "• NO claims of statistical improvement or effectiveness"
    p = tf.add_paragraph()
    p.text = "• Scenario-based analysis = conceptual exploration, NOT case studies"
    p = tf.add_paragraph()
    p.text = "• All findings are positioned as CONCEPTUAL CONTRIBUTIONS"
    p = tf.add_paragraph()
    p.text = "• Explicit limitations acknowledged throughout"
    
    # SLIDE 3: BACKGROUND & PROBLEM STATEMENT
    slide = prs.slides.add_slide(content_slide_layout)
    title = slide.shapes.title
    content = slide.placeholders[1]
    
    title.text = "Background & Problem Statement"
    tf = content.text_frame
    tf.text = "THE CONTEMPORARY EDUCATIONAL CHALLENGE"
    p = tf.add_paragraph()
    p.text = "• Rising mental health challenges among school students globally"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "  - Stress, anxiety, emotional instability (WHO, 2021)"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "  - Academic pressure and performance anxiety"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "  - Peer conflicts and social isolation"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "• Overemphasis on academic performance"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "  - Cognitive development prioritized over holistic growth"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "  - Limited integration of values and emotional intelligence"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "• Fragmented life-skills and mental-health interventions"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "  - Lack of culturally grounded frameworks"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "  - Need for comprehensive value-based approaches"
    p.level = 2
    
    # SLIDE 4: AIM & OBJECTIVES
    slide = prs.slides.add_slide(content_slide_layout)
    title = slide.shapes.title
    content = slide.placeholders[1]
    
    title.text = "Aim & Objectives"
    tf = content.text_frame
    tf.text = "Primary Aim:"
    p = tf.add_paragraph()
    p.text = "To critically review value-based education literature and analyze traditional Indian texts as conceptual frameworks for mental health support"
    p = tf.add_paragraph()
    p.text = "Specific Objectives:"
    p = tf.add_paragraph()
    p.text = "• Synthesize literature on value-based education (2015-2025)"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• Analyze traditional Indian texts as potential conceptual frameworks"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• Employ scenario-based analytical approach"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• Identify limitations and research gaps"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• Position work as conceptual contribution requiring empirical validation"
    p.level = 1
    
    # SLIDE 5: METHODOLOGY
    slide = prs.slides.add_slide(content_slide_layout)
    title = slide.shapes.title
    content = slide.placeholders[1]
    
    title.text = "Methodology"
    tf = content.text_frame
    tf.text = "Review Type: Narrative literature review (not systematic review)"
    p = tf.add_paragraph()
    p.text = "Databases: Google Scholar, PubMed, Semantic Scholar"
    p = tf.add_paragraph()
    p.text = "Time Period: 2015-2025"
    p = tf.add_paragraph()
    p.text = "Search Focus:"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• Value-based education AND mental health"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "• Traditional Indian texts AND education"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "• Narrative approaches to moral/ethical education"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "Analytical Approach: Scenario-based synthesis"
    p = tf.add_paragraph()
    p.text = "• Scenarios = analytical units, NOT case studies"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• Conceptual exploration, NOT empirical evidence"
    p.level = 1
    
    # SLIDE 6: RATIONALE FOR SCENARIO-BASED ANALYSIS
    slide = prs.slides.add_slide(content_slide_layout)
    title = slide.shapes.title
    content = slide.placeholders[1]
    
    title.text = "Rationale for Scenario-Based Analysis"
    tf = content.text_frame
    tf.text = "WHY SCENARIOS?"
    p = tf.add_paragraph()
    p.text = "• Scenarios as analytical units derived from literature"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "  - Common mental health challenges identified in research"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "  - Structured framework for conceptual exploration"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "• Appropriate for Conceptual Research"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "  - Allows exploration of potential applications"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "  - Facilitates critical analysis of limitations"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "  - Supports theoretical framework development"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "• Explicit Clarification"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "  - Scenarios are NOT case studies"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "  - Scenarios are NOT empirical evidence"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "  - Scenarios are analytical tools for conceptual research"
    p.level = 2
    
    # SLIDE 7: TRADITIONAL TEXTS
    slide = prs.slides.add_slide(content_slide_layout)
    title = slide.shapes.title
    content = slide.placeholders[1]
    
    title.text = "Traditional Texts - Conceptual Foundations"
    tf = content.text_frame
    tf.text = "Basavanna's Vachanas (12th century)"
    p = tf.add_paragraph()
    p.text = "• Kannada philosophical poetry"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• Themes: Social equality, ethical conduct, self-awareness"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• Emphasis: Inner development, empathy, moral responsibility"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Hitopadesha (Ancient Sanskrit)"
    p = tf.add_paragraph()
    p.text = "• Animal fables teaching practical wisdom"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• Themes: Decision-making, conflict resolution, interpersonal skills"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• Format: Narrative allegories"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Panchatantra (Classical collection)"
    p = tf.add_paragraph()
    p.text = "• Animal stories for life skills education"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• Themes: Cooperation, wisdom, ethical behavior"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• Approach: Allegorical teaching"
    p.level = 1
    
    # SCENARIO SLIDES (8-17)
    scenarios = [
        {
            "num": 1,
            "title": "Academic Stress & Performance Anxiety",
            "problem": "High levels of stress, sleep disturbances, anxiety disorders in competitive academic environments (Pandey, 2019; Sharma, 2018)",
            "text": "Basavanna's Vachanas - emphasis on inner peace, detachment from external validation, intrinsic motivation",
            "link": "Philosophical shift from external metrics to internal growth might support stress reduction (theoretical connection)",
            "limitation": "No empirical studies validate effectiveness. Framework requires rigorous controlled testing."
        },
        {
            "num": 2,
            "title": "Peer Conflict & Aggression",
            "problem": "Peer conflicts and aggressive behaviors prevalent in schools, contributing to emotional distress and social isolation (Joshi & Mehta, 2021)",
            "text": "Hitopadesha - fables addressing conflict resolution, understanding perspectives, finding mutually beneficial solutions",
            "link": "Narrative-based discussions might provide frameworks for navigating interpersonal challenges. Allegorical format facilitates reflection.",
            "limitation": "Transfer of narrative lessons to real-world conflict resolution requires empirical validation. Cultural factors may limit generalizability."
        },
        {
            "num": 3,
            "title": "Low Self-Esteem & Identity Confusion",
            "problem": "Adolescent identity development issues, low self-worth associated with poor academic performance and social withdrawal (Singh & Verma, 2021)",
            "text": "Basavanna's Vachanas - emphasis on inherent worth of all individuals regardless of social status, promoting self-acceptance",
            "link": "Philosophical teachings about universal human value might conceptually support positive identity formation",
            "limitation": "No studies directly examine impact on self-esteem. Historical context may require adaptation. May not resonate across all cultural backgrounds."
        },
        {
            "num": 4,
            "title": "Bullying & Social Exclusion",
            "problem": "Bullying and social exclusion represent serious mental health threats, associated with depression, anxiety, suicidal ideation (Kumaran & Singh, 2019)",
            "text": "Panchatantra - stories about consequences of harmful behavior, importance of treating others with respect, how exclusion harms perpetrators",
            "link": "Allegorical stories might provide safe spaces for discussing bullying. Narrative format could facilitate empathy development.",
            "limitation": "Narrative-based interventions for bullying require comprehensive empirical evaluation. Effectiveness in addressing contemporary bullying dynamics remains unproven."
        },
        {
            "num": 5,
            "title": "Anger & Emotional Dysregulation",
            "problem": "Emotional dysregulation, particularly anger management issues, impacts well-being and classroom dynamics. Linked to academic and social difficulties (Kumar, 2018)",
            "text": "Basavanna's Vachanas - emphasis on self-control, emotional mastery, understanding transient nature of emotions",
            "link": "Philosophical approaches to emotional awareness might conceptually support regulation skills. Aligns with cognitive-behavioral principles.",
            "limitation": "Abstract nature of philosophical teachings may not be accessible to all age groups. Direct translation to practical strategies requires empirical validation."
        },
        {
            "num": 6,
            "title": "Dishonesty & Ethical Drift",
            "problem": "Academic dishonesty and ethical lapses raise concerns about character development. Value-based interventions might support ethical decision-making (Pandey, 2019)",
            "text": "Hitopadesha and Panchatantra - stories illustrating consequences of dishonesty, value of integrity, how short-term gains lead to long-term losses",
            "link": "Moral narratives might provide frameworks for ethical reasoning. Story format makes abstract concepts more concrete.",
            "limitation": "Relationship between narrative exposure and actual ethical behavior requires empirical investigation. Cultural and generational differences may affect relevance."
        },
        {
            "num": 7,
            "title": "Digital Addiction & Attention Fragmentation",
            "problem": "Digital device overuse leads to attention difficulties, sleep problems, social isolation. Increasing technology-related mental health issues (Nair & Sharma, 2020)",
            "text": "Basavanna's Vachanas - emphasis on mindfulness, present-moment awareness, concepts related to attention regulation",
            "link": "Philosophical teachings about focus and awareness might provide frameworks for understanding attention and distraction",
            "limitation": "Historical context predates digital technology entirely. Direct application requires significant adaptation. No empirical studies examine this connection."
        },
        {
            "num": 8,
            "title": "Peer Pressure & Impulsive Decision-Making",
            "problem": "Adolescent susceptibility to peer pressure and impulsive decision-making contributes to risky behaviors and negative mental health outcomes (Sharma & Rao, 2017)",
            "text": "Hitopadesha - stories about wise decision-making, considering consequences, resisting harmful influences, importance of independent judgment",
            "link": "Stories about decision-making might provide frameworks for understanding choice, consequence, and peer influence. Narrative format could facilitate discussion of sensitive topics.",
            "limitation": "Narrative exposure alone may not translate to improved decision-making in real-world peer pressure situations. Comprehensive skill-building interventions likely necessary."
        },
        {
            "num": 9,
            "title": "Demotivation & Loss of Purpose",
            "problem": "Student demotivation and lack of purpose associated with poor academic performance and mental health issues. Meaning-making supports well-being (Narayan, 2015)",
            "text": "Basavanna's Vachanas - emphasis on finding purpose through service and ethical living. Meaning comes from contributing to others' well-being.",
            "link": "Philosophical discussions of purpose and meaning might conceptually support motivation and well-being",
            "limitation": "Abstract nature of purpose and meaning may be difficult for younger students to grasp. Spiritual context may require secular adaptation. Empirical validation lacking."
        },
        {
            "num": 10,
            "title": "Social Responsibility & Empathy Deficit",
            "problem": "Research indicates declining empathy levels and reduced engagement with social responsibility. Concerns about social cohesion and individual well-being (Desai & Patel, 2020)",
            "text": "All three sources emphasize social responsibility, compassion, concern for others' welfare. Basavanna's Vachanas particularly stress equality and service.",
            "link": "Narratives emphasizing empathy and social responsibility might provide frameworks for understanding interconnectedness and mutual care. Story-based approaches could facilitate empathy development.",
            "limitation": "Relationship between narrative exposure and actual empathetic behavior requires empirical investigation. Cultural and individual differences may affect narrative impact."
        }
    ]
    
    for scenario in scenarios:
        slide = prs.slides.add_slide(content_slide_layout)
        title = slide.shapes.title
        content = slide.placeholders[1]
        
        title.text = f"Scenario {scenario['num']} - {scenario['title']}"
        tf = content.text_frame
        tf.text = "PROBLEM:"
        p = tf.add_paragraph()
        p.text = scenario['problem']
        p.level = 1
        p = tf.add_paragraph()
        p.text = "TRADITIONAL TEXT:"
        p = tf.add_paragraph()
        p.text = scenario['text']
        p.level = 1
        p = tf.add_paragraph()
        p.text = "CONCEPTUAL LINK:"
        p = tf.add_paragraph()
        p.text = scenario['link']
        p.level = 1
        p = tf.add_paragraph()
        p.text = "LIMITATION:"
        p = tf.add_paragraph()
        p.text = scenario['limitation']
        p.level = 1
    
    # SLIDE 18: SYNTHESIS - WHAT THEY CAN SUPPORT
    slide = prs.slides.add_slide(content_slide_layout)
    title = slide.shapes.title
    content = slide.placeholders[1]
    
    title.text = "Synthesis - What Value-Based Narratives Can Conceptually Support"
    tf = content.text_frame
    tf.text = "POTENTIAL CONCEPTUAL CONTRIBUTIONS"
    p = tf.add_paragraph()
    p.text = "• Emotional Awareness: Self-reflection and emotional understanding frameworks"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• Moral Reasoning: Ethical decision-making frameworks through narratives"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• Empathy Development: Diverse perspectives and experiences through stories"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• Resilience Building: Teachings about inner strength and perseverance"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• Cultural Identity: Connection to traditional wisdom may support belonging"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "IMPORTANT: These are CONCEPTUAL possibilities, NOT proven outcomes"
    
    # SLIDE 19: SYNTHESIS - WHAT THEY CANNOT ADDRESS
    slide = prs.slides.add_slide(content_slide_layout)
    title = slide.shapes.title
    content = slide.placeholders[1]
    
    title.text = "Synthesis - What They Cannot Address"
    tf = content.text_frame
    tf.text = "CRITICAL LIMITATIONS"
    p = tf.add_paragraph()
    p.text = "Narrative approaches alone CANNOT:"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• Replace clinical interventions for serious mental health conditions"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "• Guarantee behavioral change from narrative exposure"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "• Address all cultural and individual contexts"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "• Provide immediate solutions (require time and reflection)"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "• Substitute for comprehensive mental health programs"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "Isolated narrative exposure is INSUFFICIENT"
    
    # SLIDE 20: WHERE EMPIRICAL EVIDENCE IS LACKING
    slide = prs.slides.add_slide(content_slide_layout)
    title = slide.shapes.title
    content = slide.placeholders[1]
    
    title.text = "Where Empirical Evidence is Lacking"
    tf = content.text_frame
    tf.text = "SIGNIFICANT RESEARCH GAPS"
    p = tf.add_paragraph()
    p.text = "• Intervention Studies: Limited controlled trials examining traditional text-based interventions"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• Longitudinal Research: Absence of long-term follow-up studies"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• Standardized Assessment: Lack of validated tools for measuring outcomes"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• Comparative Effectiveness: No studies comparing traditional narrative approaches to other interventions"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• Implementation Research: Limited understanding of effective delivery methods and teacher training needs"
    p.level = 1
    
    # SLIDE 21: CRITICAL LIMITATIONS
    slide = prs.slides.add_slide(content_slide_layout)
    title = slide.shapes.title
    content = slide.placeholders[1]
    
    title.text = "Critical Limitations"
    tf = content.text_frame
    tf.text = "EXPLICIT ACKNOWLEDGMENT"
    p = tf.add_paragraph()
    p.text = "Methodological Limitations:"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• Narrative review (not systematic review or meta-analysis)"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "• Potential selection and publication bias"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "• Language limitations"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "Conceptual Limitations:"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• Cultural generalization risks"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "• Temporal distance requiring adaptation"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "• Abstract nature may limit accessibility"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "• Contextual differences from historical settings"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "Implementation Challenges:"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• Teacher training requirements"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "• Curriculum integration difficulties"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "• Resource and assessment challenges"
    p.level = 2
    
    # SLIDE 22: CONCLUSION
    slide = prs.slides.add_slide(content_slide_layout)
    title = slide.shapes.title
    content = slide.placeholders[1]
    
    title.text = "Conclusion & Future Directions"
    tf = content.text_frame
    tf.text = "POSITION STATEMENT"
    p = tf.add_paragraph()
    p.text = "This work is a CONCEPTUAL CONTRIBUTION to educational psychology:"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• Does NOT claim empirical effectiveness"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "• Does NOT recommend specific interventions"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "• Provides framework for understanding potential applications"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "• Emphasizes critical need for rigorous empirical research"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "Future Research Priorities:"
    p = tf.add_paragraph()
    p.text = "1. Controlled intervention studies"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "2. Development of validated assessment tools"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "3. Mechanism research (how narratives might influence mental health)"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "4. Age-appropriate and culturally sensitive adaptations"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "5. Longitudinal impact assessment"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "6. Evidence-based implementation strategies"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Integration of traditional wisdom with contemporary needs = promising area requiring careful empirical validation"
    
    # Save presentation
    filename = "Value_Based_Education_Literature_Review_PPT.pptx"
    prs.save(filename)
    print(f"Presentation saved as {filename}")
    print(f"Total slides: {len(prs.slides)}")

if __name__ == "__main__":
    try:
        create_presentation()
        print("\n✓ PowerPoint presentation created successfully!")
        print("\nNote: Review the presentation and adjust formatting as needed.")
    except ImportError:
        print("ERROR: python-pptx library not found.")
        print("Please install it using: pip install python-pptx")
    except Exception as e:
        print(f"Error creating presentation: {e}")



